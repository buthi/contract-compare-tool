#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
合同差异对比 Flask 后端
整合 contract_diff_agent.py + json_to_html_report.py 核心逻辑
API 端点：
  POST /api/compare   - 接收两个文件，返回差异 JSON
  GET  /              - 服务前端页面
  GET  /health        - 健康检查
"""

from __future__ import annotations

import io
import json
import os
import re
import tempfile
import uuid
from dataclasses import dataclass, asdict
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from flask import Flask, jsonify, request, send_from_directory
from flask_cors import CORS
from werkzeug.utils import secure_filename

# ─────────────────────────────────────────────
# Flask 初始化
# ─────────────────────────────────────────────

BASE_DIR = Path(__file__).parent
PUBLIC_DIR = BASE_DIR / "public"
UPLOAD_DIR = BASE_DIR / "uploads"
UPLOAD_DIR.mkdir(exist_ok=True)

ALLOWED_EXTENSIONS = {"docx", "pdf", "txt", "md"}
MAX_FILE_SIZE = 20 * 1024 * 1024  # 20 MB

app = Flask(__name__, static_folder=str(PUBLIC_DIR), static_url_path="")
CORS(app)
app.config["MAX_CONTENT_LENGTH"] = MAX_FILE_SIZE


def allowed_file(filename: str) -> bool:
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


# ─────────────────────────────────────────────
# Data classes (来自 contract_diff_agent.py)
# ─────────────────────────────────────────────

@dataclass
class Clause:
    idx: int
    level: str
    clause_id: str
    title: str
    raw_text: str
    normalized_text: str
    source: str


@dataclass
class ClauseDiff:
    clause_id: str
    title: str
    status: str          # same | modified | added | removed
    similarity: float
    template_excerpt: str
    target_excerpt: str
    template_raw: str
    target_raw: str
    substantive: bool
    reason: str
    ai_checked: bool = False
    ai_material: bool = False
    ai_category: str = ""
    ai_reason: str = ""
    llm_summary: str = ""


# ─────────────────────────────────────────────
# 文本提取 (IO)
# ─────────────────────────────────────────────

def read_docx_text(path: Path) -> str:
    from docx import Document  # type: ignore
    doc = Document(str(path))
    parts: List[str] = []
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            parts.append(txt)
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
            if row_text:
                parts.append(row_text)
    return "\n".join(parts)


def read_pdf_text(path: Path) -> str:
    text_parts: List[str] = []
    try:
        import fitz  # type: ignore
        doc = fitz.open(str(path))
        for page in doc:
            txt = page.get_text("text", sort=True)
            if txt:
                text_parts.append(txt)
        if text_parts:
            return smart_join_pdf_lines("\n".join(text_parts))
    except Exception:
        pass

    try:
        import pdfplumber  # type: ignore
        with pdfplumber.open(str(path)) as pdf:
            for page in pdf.pages:
                txt = page.extract_text() or ""
                if txt:
                    text_parts.append(txt)
        if text_parts:
            return smart_join_pdf_lines("\n".join(text_parts))
    except Exception as e:
        raise RuntimeError(f"无法解析 PDF 文本: {e}")

    return ""


def read_text_file(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return read_docx_text(path)
    if suffix == ".pdf":
        return read_pdf_text(path)
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    raise ValueError(f"不支持的文件格式: {suffix}")


# ─────────────────────────────────────────────
# 清洗 & 归一化
# ─────────────────────────────────────────────

def clean_text(text: str) -> str:
    text = text.replace("\xa0", " ").replace("\u3000", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\r", "\n", text)
    text = re.sub(r"\n{2,}", "\n", text)
    return text.strip()


PDF_NOISE_PATTERNS = [
    r"^\d{1,3}$",
    r"^合同编号[:：].*$",
    r"^《联合经营协议》签署页$",
]


def is_pdf_noise_line(line: str) -> bool:
    s = line.strip()
    if not s:
        return False
    return any(re.match(p, s) for p in PDF_NOISE_PATTERNS)


ARTICLE_RE = re.compile(r"^(第[一二三四五六七八九十百零〇\d]+条)\s*(.*)$")


def is_clause_start(line: str) -> Optional[Tuple[str, str, str]]:
    s = line.strip()
    if not s:
        return None
    if re.match(r"^第[一二三四五六七八九十百零〇\d]+条", s):
        m = re.match(r"^(第[一二三四五六七八九十百零〇\d]+条)\s*(.*)$", s)
        if m:
            return ("article", m.group(1), m.group(2).strip())
    if re.match(r"^[一二三四五六七八九十]+、", s):
        m = re.match(r"^([一二三四五六七八九十]+、)\s*(.*)$", s)
        if m:
            return ("section", m.group(1), m.group(2).strip())
    if re.match(r"^（[一二三四五六七八九十]+）", s):
        m = re.match(r"^(（[一二三四五六七八九十]+）)\s*(.*)$", s)
        if m:
            return ("subsection", m.group(1), m.group(2).strip())
    if re.match(r"^\d+(?:\.\d+){0,3}[、.\s]", s):
        m = re.match(r"^(\d+(?:\.\d+){0,3})[、.\s]+(.*)$", s)
        if m:
            return ("numeric", m.group(1), m.group(2).strip())
    return None


def smart_join_pdf_lines(text: str) -> str:
    text = text.replace("\xa0", " ").replace("\u3000", " ")
    raw_lines = [re.sub(r"[ \t]+", " ", ln).strip() for ln in text.splitlines()]
    lines: List[str] = []
    for ln in raw_lines:
        if not ln:
            lines.append("")
            continue
        if is_pdf_noise_line(ln):
            continue
        lines.append(ln)

    merged: List[str] = []
    buf = ""

    def flush():
        nonlocal buf
        if buf.strip():
            merged.append(buf.strip())
        buf = ""

    for line in lines:
        if not line:
            flush()
            continue
        if is_clause_start(line):
            flush()
            buf = line
            continue
        if not buf:
            buf = line
            continue
        prev = buf.rstrip()
        if re.search(r"[。！？；]$", prev):
            flush()
            buf = line
            continue
        if re.search(r"[\u4e00-\u9fff]$", prev) or re.match(
            r"^[\u4e00-\u9fff，。；：、】【（）《》""'']", line
        ):
            sep = ""
        else:
            sep = " "
        buf = prev + sep + line

    flush()
    out = "\n".join(merged)
    out = re.sub(r"\n{3,}", "\n\n", out)
    return out.strip()


PLACEHOLDER_PATTERNS = [
    r"【[^】]{1,120}】",
    r"\[[^\]]{1,120}\]",
    r"\([^\)]{1,120}\)",
    r"（[^）]{1,120}）",
    r"＜[^＞]{1,120}＞",
    r"<[^>]{1,120}>",
]

VARIABLE_PATTERNS = [
    r"\b\d{4}[年/\-.]\d{1,2}[月/\-.]\d{1,2}日?\b",
    r"\b\d{1,2}[/:]\d{1,2}(?::\d{1,2})?\b",
    r"¥\s*\d+(?:,\d{3})*(?:\.\d+)?",
    r"人民币\s*\d+(?:,\d{3})*(?:\.\d+)?\s*元?",
    r"\b\d+(?:,\d{3})*(?:\.\d+)?\s*元\b",
    r"\b\d+(?:\.\d+)?%\b",
    r"\b\d{11,19}\b",
    r"\b1[3-9]\d{9}\b",
    r"\b[A-Z0-9]{8,32}\b",
]


def normalize_for_diff(text: str) -> str:
    out = clean_text(text)
    for pat in PLACEHOLDER_PATTERNS:
        out = re.sub(pat, "<PLACEHOLDER>", out)
    out = re.sub(r"_{3,}|＿{3,}|﹍{3,}|\.{5,}|·{3,}|x{3,}|X{3,}", "<PLACEHOLDER>", out)
    for pat in VARIABLE_PATTERNS:
        out = re.sub(pat, "<VAR>", out)
    out = out.replace("：", ":").replace("；", ";").replace("，", ",").replace("。", ".")
    out = re.sub(r"\s+", " ", out)
    return out.strip()


# ─────────────────────────────────────────────
# 条款切分
# ─────────────────────────────────────────────

def split_into_articles(text: str, source: str) -> List[Clause]:
    text = clean_text(text)
    lines = [ln.strip() for ln in text.split("\n") if ln.strip()]

    clauses: List[Clause] = []
    current_id = "PREAMBLE"
    current_title = "前言"
    current_level = "preamble"
    current_lines: List[str] = []
    idx = 0

    def flush():
        nonlocal idx, current_id, current_title, current_level, current_lines
        if not current_lines:
            return
        raw = "\n".join(current_lines).strip()
        clauses.append(
            Clause(
                idx=idx,
                level=current_level,
                clause_id=current_id,
                title=current_title,
                raw_text=raw,
                normalized_text=normalize_for_diff(raw),
                source=source,
            )
        )
        idx += 1
        current_lines = []

    for line in lines:
        m = ARTICLE_RE.match(line)
        if m:
            flush()
            current_id = m.group(1)
            current_title = m.group(2).strip() or m.group(1)
            current_level = "article"
            current_lines = [line]
        else:
            current_lines.append(line)

    flush()

    if not clauses:
        whole = clean_text(text)
        clauses = [
            Clause(
                idx=0,
                level="whole",
                clause_id="WHOLE_DOC",
                title="全文",
                raw_text=whole,
                normalized_text=normalize_for_diff(whole),
                source=source,
            )
        ]
    return clauses


def _field(obj, name, default=""):
    if isinstance(obj, dict):
        return obj.get(name, default)
    return getattr(obj, name, default)


def regroup_target_to_article_level(target_clauses: List[Clause]) -> List[Clause]:
    regrouped: List[Clause] = []
    current: Optional[Clause] = None

    for item in target_clauses:
        clause_id = str(_field(item, "clause_id", "")).strip()
        title = str(_field(item, "title", "")).strip()
        raw_text = str(_field(item, "raw_text", "")).strip()
        normalized_text = str(_field(item, "normalized_text", "")).strip()
        level = str(_field(item, "level", "")).strip()
        source = str(_field(item, "source", "target")).strip() or "target"

        if clause_id == "PREAMBLE":
            regrouped.append(
                Clause(
                    idx=len(regrouped),
                    level="preamble",
                    clause_id="PREAMBLE",
                    title="前言",
                    raw_text=raw_text,
                    normalized_text=normalized_text,
                    source=source,
                )
            )
            current = None
            continue

        if ARTICLE_RE.match(clause_id) or ARTICLE_RE.match(title):
            if current is not None:
                regrouped.append(current)
            article_name = clause_id if ARTICLE_RE.match(clause_id) else title
            current = Clause(
                idx=-1,
                level="article",
                clause_id=article_name,
                title=title if title else article_name,
                raw_text=raw_text,
                normalized_text=normalized_text,
                source=source,
            )
            continue

        if current is not None:
            if raw_text:
                current.raw_text = (current.raw_text + "\n" + raw_text).strip()
            if normalized_text:
                current.normalized_text = (current.normalized_text + "\n" + normalized_text).strip()
        else:
            if regrouped and regrouped[-1].clause_id == "PREAMBLE":
                if raw_text:
                    regrouped[-1].raw_text = (regrouped[-1].raw_text + "\n" + raw_text).strip()
                if normalized_text:
                    regrouped[-1].normalized_text = (
                        regrouped[-1].normalized_text + "\n" + normalized_text
                    ).strip()
            else:
                regrouped.append(
                    Clause(
                        idx=len(regrouped),
                        level=level or "other",
                        clause_id=clause_id,
                        title=title,
                        raw_text=raw_text,
                        normalized_text=normalized_text,
                        source=source,
                    )
                )

    if current is not None:
        regrouped.append(current)

    for i, item in enumerate(regrouped):
        item.idx = i

    return regrouped


# ─────────────────────────────────────────────
# 对齐 & 规则差异判断
# ─────────────────────────────────────────────

def clause_key(c: Clause) -> str:
    title = re.sub(r"\s+", "", c.title)
    return f"{c.clause_id}::{title}".strip()


def excerpt(text: str, limit: int = 300) -> str:
    text = clean_text(text).replace("\n", " ")
    if len(text) <= limit:
        return text
    return text[:limit] + " ..."


def similarity(a: str, b: str) -> float:
    return SequenceMatcher(None, a, b).ratio()


def substantive_diff_reason(
    template_clause: Clause, target_clause: Clause, sim: float
) -> Tuple[bool, str]:
    a = template_clause.normalized_text
    b = target_clause.normalized_text

    if a == b:
        return False, "normalized texts identical; likely only placeholders/variables changed"
    if sim >= 0.985:
        return False, f"very high normalized similarity ({sim:.3f}); likely minor wording only"

    len_a = len(a)
    len_b = len(b)
    length_ratio = min(len_a, len_b) / max(len_a, len_b) if max(len_a, len_b) else 1.0

    if sim < 0.85:
        return True, f"normalized similarity is low ({sim:.3f}); possible substantive rewrite"
    if length_ratio < 0.72:
        return True, f"substantial length change detected (ratio={length_ratio:.3f})"

    trigger_words = ["违约", "赔偿", "解除", "终止", "保密", "争议", "仲裁", "责任", "期限", "费用", "税", "退款", "违约金"]
    trig_a = {w for w in trigger_words if w in a}
    trig_b = {w for w in trigger_words if w in b}
    if trig_a != trig_b:
        return True, f"legal trigger words changed: template={sorted(trig_a)}, target={sorted(trig_b)}"

    if sim < 0.93:
        return True, f"moderate normalized difference remains ({sim:.3f}) after placeholder filtering"

    return False, f"difference limited after normalization (similarity={sim:.3f})"


def align_clauses(template_clauses: List[Clause], target_clauses: List[Clause]) -> List[ClauseDiff]:
    diffs: List[ClauseDiff] = []
    target_by_key: Dict[str, Clause] = {clause_key(c): c for c in target_clauses}
    used_target_keys: set = set()

    for tc in template_clauses:
        key = clause_key(tc)
        matched = target_by_key.get(key)

        if matched:
            used_target_keys.add(key)
            sim = similarity(tc.normalized_text, matched.normalized_text)
            substantive, reason = substantive_diff_reason(tc, matched, sim)
            diffs.append(
                ClauseDiff(
                    clause_id=tc.clause_id,
                    title=tc.title,
                    status="same" if (not substantive and tc.normalized_text == matched.normalized_text) else "modified",
                    similarity=sim,
                    template_excerpt=excerpt(tc.raw_text),
                    target_excerpt=excerpt(matched.raw_text),
                    template_raw=tc.raw_text,
                    target_raw=matched.raw_text,
                    substantive=substantive,
                    reason=reason,
                )
            )
        else:
            candidates = [c for c in target_clauses if c.clause_id == tc.clause_id] or target_clauses
            best: Optional[Clause] = None
            best_score = -1.0
            for cand in candidates:
                score = similarity(tc.normalized_text, cand.normalized_text)
                if tc.title and cand.title:
                    score = max(score, 0.35 * similarity(tc.title, cand.title) + 0.65 * score)
                if score > best_score:
                    best_score = score
                    best = cand

            if best is not None and best_score >= 0.78 and clause_key(best) not in used_target_keys:
                used_target_keys.add(clause_key(best))
                substantive, reason = substantive_diff_reason(tc, best, best_score)
                diffs.append(
                    ClauseDiff(
                        clause_id=tc.clause_id,
                        title=tc.title or best.title,
                        status="modified",
                        similarity=best_score,
                        template_excerpt=excerpt(tc.raw_text),
                        target_excerpt=excerpt(best.raw_text),
                        template_raw=tc.raw_text,
                        target_raw=best.raw_text,
                        substantive=substantive,
                        reason=reason + " (fuzzy-aligned)",
                    )
                )
            else:
                diffs.append(
                    ClauseDiff(
                        clause_id=tc.clause_id,
                        title=tc.title,
                        status="removed",
                        similarity=0.0,
                        template_excerpt=excerpt(tc.raw_text),
                        target_excerpt="",
                        template_raw=tc.raw_text,
                        target_raw="",
                        substantive=True,
                        reason="template clause not found in target",
                    )
                )

    for tgt in target_clauses:
        key = clause_key(tgt)
        if key not in used_target_keys:
            diffs.append(
                ClauseDiff(
                    clause_id=tgt.clause_id,
                    title=tgt.title,
                    status="added",
                    similarity=0.0,
                    template_excerpt="",
                    target_excerpt=excerpt(tgt.raw_text),
                    template_raw="",
                    target_raw=tgt.raw_text,
                    substantive=True,
                    reason="target contains an unmatched clause",
                )
            )

    return diffs


# ─────────────────────────────────────────────
# AI 判断（可选，需要配置 LLM）
# ─────────────────────────────────────────────

def judge_diffs_with_llm(
    diffs: List[ClauseDiff],
    base_url: str,
    api_key: str,
    model: str,
    temperature: float = 0.1,
) -> None:
    import requests as req_lib

    url = base_url.rstrip("/") + "/chat/completions"
    headers = {
        "Content-Type": "application/json",
        "Accept": "application/json",
        "Authorization": f"Bearer {api_key.strip()}",
    }

    for d in diffs:
        if d.status == "same" and not d.substantive:
            continue
        user_prompt = f"""你是一名资深合同审查助手。请判断以下两段条款之间是否存在【法律效果上的重大差异】。

核心判断原则：
- 只有差异足以改变合同一方或多方的法律地位、责任范围、履行要求、收益分配、违约后果或救济方式时，才认定为【重大差异】。
- 企业名称、日期、金额、账号、占位符替换、格式调整等通常不属于重大差异。
- 权利义务变化、违约责任变化、合同期限/解除条件变化等通常属于重大差异。

请输出严格 JSON，不含额外解释：
{{
  "material": true/false,
  "category": "payment|liability|termination|dispute|confidentiality|obligation|rights|procedure|other",
  "reason": "一句话说明是否属于重大差异（聚焦法律效果）",
  "summary": "1-3句话总结核心差异"
}}

条款编号：{d.clause_id}
条款标题：{d.title}
状态：{d.status}

模板条款：
{d.template_raw or "[无]"}

目标条款：
{d.target_raw or "[无]"}
""".strip()

        payload = {
            "model": model,
            "messages": [
                {"role": "system", "content": "你是严谨的法务条款差异审查助手，只能输出 JSON。"},
                {"role": "user", "content": user_prompt},
            ],
            "temperature": temperature,
        }

        try:
            resp = req_lib.post(url, headers=headers, json=payload, timeout=180)
            d.ai_checked = True
            if resp.status_code != 200:
                d.ai_reason = f"LLM 调用失败: HTTP {resp.status_code}"
                continue
            data = resp.json()
            content = data["choices"][0]["message"]["content"].strip()
            content = re.sub(r"^```json\s*", "", content)
            content = re.sub(r"^```\s*", "", content)
            content = re.sub(r"\s*```$", "", content)
            obj = json.loads(content)
            d.ai_material = bool(obj.get("material", False))
            d.ai_category = str(obj.get("category", "")).strip()
            d.ai_reason = str(obj.get("reason", "")).strip()
            d.llm_summary = str(obj.get("summary", "")).strip()
        except Exception as e:
            d.ai_checked = True
            d.ai_reason = f"LLM 解析失败: {e}"


# ─────────────────────────────────────────────
# 数据归一化（来自 json_to_html_report.py）
# ─────────────────────────────────────────────

def normalize_data(
    diffs: List[ClauseDiff],
    template_file: str = "",
    target_file: str = "",
) -> Dict[str, Any]:
    diff_dicts = [asdict(d) for d in diffs]
    material_count = sum(
        1 for d in diffs
        if d.ai_material or (not d.ai_checked and d.substantive)
    )
    return {
        "meta": {
            "templateFile": template_file or "未提供",
            "targetFile": target_file or "未提供",
            "totalDiffs": len(diff_dicts),
            "materialDiffs": material_count,
            "source": "在线比对",
        },
        "diffs": diff_dicts,
    }


# ─────────────────────────────────────────────
# Flask 路由
# ─────────────────────────────────────────────

@app.route("/health")
def health():
    return jsonify({"ok": True, "service": "contract-diff-api"})


@app.route("/")
def index():
    return send_from_directory(str(PUBLIC_DIR), "contract.html")


@app.route("/api/compare", methods=["POST"])
def api_compare():
    """
    接收 multipart/form-data:
      - template: 标准文档文件
      - target:   对比文档文件
      - llm_base_url: (可选) OpenAI 兼容接口 base url
      - llm_api_key:  (可选) API key
      - llm_model:    (可选) 模型名称

    返回 JSON:
      {
        "ok": true,
        "meta": {...},
        "diffs": [...]
      }
    """
    if "template" not in request.files or "target" not in request.files:
        return jsonify({"ok": False, "error": "请同时上传 template（标准文档）和 target（对比文档）"}), 400

    template_file = request.files["template"]
    target_file = request.files["target"]

    if not template_file.filename or not target_file.filename:
        return jsonify({"ok": False, "error": "文件名为空"}), 400

    if not allowed_file(template_file.filename):
        return jsonify({"ok": False, "error": f"标准文档格式不支持，仅支持: {', '.join(ALLOWED_EXTENSIONS)}"}), 400

    if not allowed_file(target_file.filename):
        return jsonify({"ok": False, "error": f"对比文档格式不支持，仅支持: {', '.join(ALLOWED_EXTENSIONS)}"}), 400

    # 获取可选 LLM 配置
    llm_base_url = request.form.get("llm_base_url", "").strip()
    llm_api_key = request.form.get("llm_api_key", "").strip()
    llm_model = request.form.get("llm_model", "").strip()

    # 保存上传文件到临时目录
    tmp_dir = Path(tempfile.mkdtemp(dir=str(UPLOAD_DIR)))
    try:
        tpl_suffix = "." + template_file.filename.rsplit(".", 1)[-1].lower()
        tgt_suffix = "." + target_file.filename.rsplit(".", 1)[-1].lower()
        tpl_path = tmp_dir / f"template{tpl_suffix}"
        tgt_path = tmp_dir / f"target{tgt_suffix}"

        template_file.save(str(tpl_path))
        target_file.save(str(tgt_path))

        # 提取文本
        try:
            template_text = read_text_file(tpl_path)
        except Exception as e:
            return jsonify({"ok": False, "error": f"标准文档解析失败: {e}"}), 422

        try:
            target_text = read_text_file(tgt_path)
        except Exception as e:
            return jsonify({"ok": False, "error": f"对比文档解析失败: {e}"}), 422

        if len(template_text.strip()) < 10:
            return jsonify({"ok": False, "error": "标准文档内容为空或无法提取文本"}), 422
        if len(target_text.strip()) < 10:
            return jsonify({"ok": False, "error": "对比文档内容为空或无法提取文本"}), 422

        # 条款切分
        template_clauses = split_into_articles(template_text, source="template")
        target_clauses = split_into_articles(target_text, source="target")
        target_clauses = regroup_target_to_article_level(target_clauses)

        # 差异对齐
        diffs = align_clauses(template_clauses, target_clauses)

        # AI 判断（可选）
        if llm_base_url and llm_api_key and llm_model:
            try:
                judge_diffs_with_llm(
                    diffs=diffs,
                    base_url=llm_base_url,
                    api_key=llm_api_key,
                    model=llm_model,
                )
            except Exception as e:
                # AI 失败不影响主流程，仅记录
                app.logger.warning(f"AI 判断失败（已跳过）: {e}")

        # 组装返回数据
        result = normalize_data(
            diffs=diffs,
            template_file=template_file.filename,
            target_file=target_file.filename,
        )
        result["ok"] = True
        return jsonify(result)

    finally:
        # 清理临时文件
        import shutil
        try:
            shutil.rmtree(str(tmp_dir), ignore_errors=True)
        except Exception:
            pass


# ─────────────────────────────────────────────
# 主入口
# ─────────────────────────────────────────────

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    print(f"🚀 合同差异对比服务启动于 http://0.0.0.0:{port}")
    app.run(host="0.0.0.0", port=port, debug=False)
